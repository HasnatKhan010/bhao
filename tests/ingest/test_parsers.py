"""Track A tests: parsers, normalisation, validation and revision-aware publishing.

The synthetic annex mimics the real geometry (verified by the 2026-08-27 spike):
city headers as `Name (NN)` on merged triples at cols D,G,J,... , item rows with
A=sr, B=name, C=unit, then MIN/AVG/MAX per city, blocks stacked vertically at
ARBITRARY rows (the parser must never hardcode them).
"""

from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd
import pytest
import openpyxl

from ingest import validate as vld
from ingest.normalise import coerce_number, parse_unit, resolve_item
from ingest.parse_annex import parse_annex
from ingest.parse_spi import parse_spi
from ingest.publish import append_week, annex_to_frame, ItemResolver

CITY_CODES = ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10",
              "11", "12", "13", "14", "15", "16", "17"]
CITY_NAMES = ["Islamabad", "Rawalpindi", "Gujranwala", "Sialkot", "Lahore",
              "Faisalabad", "Sargodha", "Multan", "Bahawalpur", "Karachi",
              "Hyderabad", "Sukkur", "Larkana", "Peshawar", "Bannu", "Quetta",
              "Khuzdar"]
ITEMS = [
    (1, "Wheat Flour Bag", "20 Kg"),
    (2, "Onions", "1 Kg"),
    (3, "Petrol Super", "Per Litre"),
    (4, "Match Box", "Each"),
    (5, "Rice IRRI-6/9 (Sindh/Punjab)", "1 Kg"),
    (6, "Chicken Farm Broiler (Live)", "1 Kg"),
]
BLOCK_ROWS = [2, 15, 28]  # deliberately NOT the real 3/61/119


def _build_annex_xlsx(with_violation: bool = False, with_dash: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Appendix-A"
    cols = [4, 7, 10, 13, 16, 19, 22]  # D,G,J,M,P,S,V
    blocks = [(0, 7), (7, 14), (14, 17)]  # city index ranges per block
    for (start, end), head_row in zip(blocks, BLOCK_ROWS):
        for k, ci in enumerate(range(start, end)):
            col = cols[k]
            cell = ws.cell(row=head_row, column=col, value=f"{CITY_NAMES[ci]} ({CITY_CODES[ci]})")
            ws.merge_cells(start_row=head_row, start_column=col, end_row=head_row, end_column=col + 2)
        sub = ws.cell(row=head_row + 1, column=1, value="NO.")
        ws.cell(row=head_row + 1, column=2, value="DESCRIPTION")
        ws.cell(row=head_row + 1, column=3, value="UNIT")
        for j, w in enumerate(["MIN", "AVG", "MAX"] * 7):
            ws.cell(row=head_row + 1, column=4 + j, value=w)
        ws.cell(row=head_row + 2, column=1, value="1")
        ws.cell(row=head_row + 3, column=4, value=f"PRICES ON 27-08-2026")
        r = head_row + 4
        for sr, name, unit in ITEMS:
            ws.cell(row=r, column=1, value=sr)
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=3, value=unit)
            for k, ci in enumerate(range(start, end)):
                base = cols[k]
                avg = 100.0 + 10 * sr + ci
                if with_violation and sr == 2 and ci == 0:
                    ws.cell(row=r, column=base, value=999)      # min > avg
                    ws.cell(row=r, column=base + 1, value=100.0)
                    ws.cell(row=r, column=base + 2, value=110.0)
                    r += 1
                    continue
                if with_dash and sr == 4 and ci == 3:
                    ws.cell(row=r, column=base, value="-")       # not surveyed
                    ws.cell(row=r, column=base + 1, value="-")
                    ws.cell(row=r, column=base + 2, value="-")
                elif sr == 1:
                    ws.cell(row=r, column=base, value="2,200")   # string numbers
                    ws.cell(row=r, column=base + 1, value=" 2,384.15 ")
                    ws.cell(row=r, column=base + 2, value="2800")
                else:
                    ws.cell(row=r, column=base, value=round(avg - 2, 2))
                    ws.cell(row=r, column=base + 1, value=round(avg, 2))
                    ws.cell(row=r, column=base + 2, value=round(avg + 2, 2))
            r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_spi_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    p1 = wb.active
    p1.title = "Page 1"
    p1["B19"] = "SPI (51 items) 2015-16=100"
    p1["B20"] = "Expenditure Group (Quintile)"
    for i, q in enumerate(["Q1 (Upto Rs. 17,732)", "Q2", "Q3", "Q4", "Q5"]):
        r = 21 + i
        p1[f"B{r}"] = q
        p1[f"D{r}"] = "340.00"
    r = 26
    p1[f"B{r}"] = "   "
    p1[f"D{r}"] = "361.07"
    p1[f"E{r}"] = "360.90"
    p1[f"F{r}"] = "331.14"
    p1[f"G{r}"] = "0.047104"
    p1[f"H{r}"] = "9.038473"

    p2 = wb.create_sheet("Page 2")
    # section 1: risers
    rows = [
        (1, "Petrol Super", "Per Litre", "345.14", "339.33", "265.82", "1.71", "29.84"),
        (2, "Onions", "1 Kg", "166.62", "171.55", "73.78", "-2.87", "125.86"),
        (3, "Wheat Flour Bag", "20 Kg", "2658.24", "2648.00", "1830.00", "0.39", "45.28"),
    ]
    r = 7
    for sr, name, unit, p, pp, ply, wow, yoy in rows:
        p2[f"B{r}"] = sr
        p2[f"C{r}"] = name
        p2[f"D{r}"] = unit
        p2[f"E{r}"] = p
        p2[f"F{r}"] = pp
        p2[f"G{r}"] = ply
        p2[f"H{r}"] = wow
        p2[f"I{r}"] = yoy
        r += 1
    # section header between sections (must be skipped)
    p2[f"C{r}"] = "ii. Items whose prices decreased"
    r += 1
    for sr, name, unit, p, pp, ply, wow, yoy in rows:
        p2[f"B{r}"] = sr
        p2[f"C{r}"] = name
        p2[f"D{r}"] = unit
        p2[f"E{r}"] = p
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def annex_bytes():
    return _build_annex_xlsx()


class TestParseAnnex:
    def test_finds_17_cities_by_regex(self, annex_bytes):
        out = parse_annex(io.BytesIO(annex_bytes), city_band=(15, 20), item_band=(4, 80))
        assert [c.code for c in out.cities] == CITY_CODES
        assert out.week_ending == dt.date(2026, 8, 27)

    def test_block_rows_not_hardcoded(self, annex_bytes):
        out = parse_annex(io.BytesIO(annex_bytes), city_band=(15, 20), item_band=(4, 80))
        rows = sorted({c.row for c in out.cities})
        assert rows == BLOCK_ROWS

    def test_item_cells(self, annex_bytes):
        out = parse_annex(io.BytesIO(annex_bytes), city_band=(15, 20), item_band=(4, 80))
        assert len({(r.sr, r.item_raw) for r in out.rows}) == len(ITEMS)
        assert len(out.rows) == len(ITEMS) * 17
        # string numbers coerced
        flour = [r for r in out.rows if r.sr == 1 and r.city_code == "01"][0]
        assert flour.price_avg == 2384.15
        # '-' → None, row kept
        match = [r for r in out.rows if r.sr == 4 and r.city_code == "04"][0]
        assert match.price_avg is None and match.price_min is None

    def test_min_max_violation_fails_loudly(self):
        bad = _build_annex_xlsx(with_violation=True)
        with pytest.raises(AssertionError, match="min<=avg<=max"):
            parse_annex(io.BytesIO(bad), city_band=(15, 20), item_band=(4, 80))

    def test_too_few_cities_fails(self, annex_bytes):
        with pytest.raises(AssertionError, match="city headers"):
            parse_annex(io.BytesIO(annex_bytes), city_band=(20, 25), item_band=(4, 80))


class TestParseSPI:
    def test_items_and_headline(self):
        out = parse_spi(io.BytesIO(_build_spi_xlsx()), min_items=5)
        assert len(out.rows) == 7  # 6 items + headline
        assert out.headline.price_this_week == 361.07
        assert out.headline.price_prev_week == 360.90
        assert out.headline.pct_change_yoy == 9.038473
        onions = [r for r in out.rows if r.item_raw == "Onions"][0]
        assert onions.price_this_week == 166.62
        assert onions.pct_change_yoy == 125.86  # PBS's own YoY, kept verbatim


class TestNormalise:
    def test_coerce_number(self):
        assert coerce_number("171.55") == 171.55
        assert coerce_number("1,171.55") == 1171.55
        assert coerce_number(" 220 ") == 220.0
        assert coerce_number("+9.00%") == 9.0
        assert coerce_number("-") is None
        assert coerce_number("") is None
        assert coerce_number("N/A") is None
        assert coerce_number(None) is None
        assert coerce_number(float("nan")) is None

    def test_parse_unit(self):
        assert parse_unit("20 Kg") == ("kg", 20.0)
        assert parse_unit("1 Kg") == ("kg", 1.0)
        assert parse_unit("Per Litre") == ("litre", 1.0)
        assert parse_unit("Per Plate") == ("plate", 1.0)
        assert parse_unit("MMBTU") == ("mmbtu", 1.0)
        assert parse_unit("1 Dozen") == ("dozen", 1.0)
        assert parse_unit("Per Minute") == ("minute", 1.0)
        assert parse_unit("40 Kg") == ("kg", 40.0)
        assert parse_unit("Per Unit") == ("unit", 1.0)
        assert parse_unit("1 mtr") == ("metre", 1.0)
        assert parse_unit("Pair") == ("pair", 1.0)
        assert parse_unit("Each") == ("each", 1.0)
        with pytest.raises(ValueError, match="unmatched unit"):
            parse_unit("25 Bangles")

    def test_resolve_item_known(self):
        assert resolve_item("Onions") == "019"
        assert resolve_item("  Chicken   Farm Broiler (Live) ") == "014"
        assert resolve_item("Rice Irri-6 (Export Quality)") == "004"  # rename pathology alias

    def test_resolve_item_unknown_raises(self):
        with pytest.raises(KeyError, match="unmatched item"):
            resolve_item("Completely New Commodity")


class TestAppendWeek:
    def _frame(self, week, avg=100.0):
        return pd.DataFrame([{
            "week_ending": week, "city_code": "05", "city_en": "Lahore", "city_ur": "لاہور",
            "item_code": "019", "item_en": "Onions", "item_ur": "پیاز",
            "unit_raw": "1 Kg", "unit_norm": "kg", "qty_norm": 1.0,
            "price_min": avg - 2, "price_avg": avg, "price_max": avg + 2,
            "price_per_unit": avg, "source": "pbs_spi_annex",
            "source_url": "https://x.example/a.xlsx",
            "ingested_at": pd.Timestamp("2026-01-10T06:00:00Z"), "revision": 0,
        }]).astype({"qty_norm": "float64", "revision": "int32",
                    "ingested_at": "datetime64[us, UTC]"})

    def test_new_week_is_revision_zero(self):
        w1, w2 = dt.date(2026, 1, 1), dt.date(2026, 1, 8)
        panel, _ = append_week(self._frame(w1, 100.0), self._frame(w2, 101.0))
        assert len(panel) == 2
        assert set(panel["revision"]) == {0}

    def test_changed_value_appends_revision_never_edits(self):
        w1 = dt.date(2026, 1, 1)
        panel, _ = append_week(self._frame(w1, 100.0), self._frame(w1, 103.0))
        assert len(panel) == 2
        rev0 = panel[panel["revision"] == 0]
        rev1 = panel[panel["revision"] == 1]
        assert float(rev0["price_avg"].iloc[0]) == 100.0, "revision 0 must never be edited"
        assert float(rev1["price_avg"].iloc[0]) == 103.0

    def test_identical_week_is_idempotent(self):
        w1 = dt.date(2026, 1, 1)
        panel, n = append_week(self._frame(w1, 100.0), self._frame(w1, 100.0))
        assert n == 0 and len(panel) == 1

    def test_tiny_difference_is_not_a_revision(self):
        w1 = dt.date(2026, 1, 1)
        panel, n = append_week(self._frame(w1, 100.0), self._frame(w1, 100.005))
        assert n == 0 and len(panel) == 1


class TestValidate:
    def _panel(self, weeks):
        frames = []
        for w in weeks:
            for city in ("05", "10"):
                frames.append({
                    "week_ending": [w], "city_code": [city], "item_code": ["019"],
                    "price_avg": [100.0], "revision": [0],
                })
        return pd.DataFrame([{k: v for row in frames for k, v in row.items()}]) if not frames else pd.concat(
            [pd.DataFrame(f) for f in frames], ignore_index=True)

    def test_week_gap_requires_seven_days(self):
        panel = self._panel([dt.date(2026, 1, 1)])
        with pytest.raises(vld.ValidationError, match="week gap"):
            vld.check_week_gap(dt.date(2026, 1, 20), panel)
        vld.check_week_gap(dt.date(2026, 1, 8), panel)  # exactly 7 days: ok

    def test_backwards_week_rejected(self):
        panel = self._panel([dt.date(2026, 1, 8)])
        with pytest.raises(vld.ValidationError, match="before"):
            vld.check_week_gap(dt.date(2026, 1, 1), panel)

    def test_coverage_collapse_rejected(self):
        panel = self._panel([dt.date(2026, 1, 1)])
        sparse = pd.DataFrame([{
            "week_ending": dt.date(2026, 1, 8), "city_code": ["05"], "item_code": ["019"],
            "price_avg": [np.nan], "revision": 0,
        }])
        with pytest.raises(vld.ValidationError, match="coverage collapse"):
            vld.check_coverage(sparse, panel)

    def test_national_crosscheck_catches_offset(self):
        city_means = {"019": 100.0, "020": 95.0}
        national = {"019": 95.0, "020": 200.0}  # tomato price under onion code
        with pytest.raises(vld.ValidationError, match="cross-check FAILED"):
            vld.national_crosscheck(city_means, national)
        vld.national_crosscheck({"019": 100.0}, {"019": 110.0})  # within 30%: ok


class TestRealFileSmoke:
    """Runs only if the real cached annex from the 2026-08-27 fetch exists."""

    def test_real_annex_parses(self):
        import glob

        paths = glob.glob("data/raw/*/*.xlsx")
        if not paths:
            pytest.skip("no cached real annex (run make spike first)")
        out = None
        for p in paths:
            try:
                cand = parse_annex(p)
            except (AssertionError, KeyError):  # not a weekly annex / error page
                continue
            if len(cand.city_codes) == 17:
                out = cand
                break
        if out is None:
            pytest.skip("no real annex among cached files")
        assert len(out.city_codes) == 17
        assert out.week_ending is not None
        assert len(out.rows) == 867
