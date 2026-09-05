"""parse_spi — the national table.

Verified geometry (spike, week 2026-08-27):
- `Page 2` holds the 51 items sorted by %change WoW in three sections (risers,
  fallers, flat — Sr restarts each section, so Sr is informational only):
  col B=Sr, C=Items, D=Units, E=price this week, F=prev week, G=same week last
  year, H=%change WoW, I=%change YoY.
- `Page 1` holds the headline: the SPI (51 items) table with one row per
  expenditure quintile plus **Combined** — that row is the headline index
  (contract item_code "000").

PBS's own WoW/YoY columns are kept verbatim so Bhao's arithmetic can be checked
against theirs.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

from ingest.normalise import coerce_number


@dataclass
class NationalRow:
    sr: int
    item_raw: str
    unit_raw: str
    price_this_week: float | None
    price_prev_week: float | None
    price_same_week_last_year: float | None
    pct_change_wow: float | None
    pct_change_yoy: float | None
    is_headline: bool = False


@dataclass
class ParsedSPI:
    rows: list[NationalRow] = field(default_factory=list)
    headline: NationalRow | None = None


def _parse_page2_items(ws) -> list[NationalRow]:
    rows: list[NationalRow] = []
    for row in ws.iter_rows(min_col=1, max_col=12):
        if len(row) < 9:
            continue
        sr_v, name_v, unit_v = row[1].value, row[2].value, row[3].value
        sr = coerce_number(sr_v)
        if sr is None or not isinstance(name_v, str) or not name_v.strip():
            continue
        name = name_v.strip()
        unit = unit_v.strip() if isinstance(unit_v, str) else ""
        p_now = coerce_number(row[4].value)
        if p_now is None:
            continue  # section headers carry no price
        rows.append(
            NationalRow(
                sr=int(sr),
                item_raw=name,
                unit_raw=unit,
                price_this_week=p_now,
                price_prev_week=coerce_number(row[5].value),
                price_same_week_last_year=coerce_number(row[6].value),
                pct_change_wow=coerce_number(row[7].value),
                pct_change_yoy=coerce_number(row[8].value),
            )
        )
    return rows


_COMBINED_RE = re.compile(r"combined", re.I)
_Q_RE = re.compile(r"^\s*Q\d\b")


def _parse_page1_headline(ws) -> NationalRow | None:
    """The Combined row of the SPI (51 items) quintile table.

    Verified geometry: B=group label (Q1..Q5 rows; the Combined row's label cell is
    BLANK — merged-cell artifact), D=this week, E=prev week, F=same week last year,
    G=WoW, H=YoY. So: find the Q1..Q5 block, then the first row after it whose
    col-D value is numeric and whose label is blank/whitespace = Combined.
    """
    rows = list(enumerate(ws.iter_rows(min_col=1, max_col=8), 1))
    q_rows = [i for i, r in rows if isinstance(r[1].value, str) and _Q_RE.match(r[1].value)]
    if not q_rows:
        return None
    last_q = q_rows[-1]
    for i, row in rows:
        if i <= last_q or i > last_q + 4:
            continue
        label = row[1].value
        p_now = coerce_number(row[3].value)  # col D
        if p_now is None:
            continue
        if isinstance(label, str) and label.strip():
            if _COMBINED_RE.search(label):  # some weeks print the label
                pass
            else:
                continue  # a different table started
        if label is not None and not isinstance(label, str):
            continue
        return NationalRow(
            sr=0,
            item_raw="SPI (Combined, 51 items)",
            unit_raw="Index",
            price_this_week=p_now,
            price_prev_week=coerce_number(row[4].value),
            price_same_week_last_year=coerce_number(row[5].value),
            pct_change_wow=coerce_number(row[6].value),
            pct_change_yoy=coerce_number(row[7].value),
            is_headline=True,
        )
    return None


def parse_spi(path, page1: str = "Page 1", page2: str = "Page 2", min_items: int = 40) -> ParsedSPI:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out = ParsedSPI()
        if page2 in wb.sheetnames:
            out.rows = _parse_page2_items(wb[page2])
        if page1 in wb.sheetnames:
            out.headline = _parse_page1_headline(wb[page1])

        assert (
            len(out.rows) >= min_items
        ), f"only {len(out.rows)} national item rows parsed from {page2}"
        assert out.headline is not None, f"no 'Combined' headline row found on {page1}"
        # the headline travels with the rows so publish sees item "000"
        out.rows.append(out.headline)
        return out
    finally:
        wb.close()
